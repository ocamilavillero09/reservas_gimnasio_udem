#!/usr/bin/env python
"""
Métricas y matriz de trazabilidad de los 18 requisitos del alcance.

Mide los requisitos de los tres integrantes con el mismo método de
metricas_camila.py (20 peticiones tras 3 de calentamiento, sobre la vista real,
verificando que la respuesta sea 2xx) y escribe "Metricas nuevas.xlsx" con la
misma estructura que "Metricas antiguas.xlsx": una hoja por requisito y la hoja
"Matriz de trazabilidad".

Además de lo que mide en vivo, calcula:
  - V(G) por función, con radon;
  - complejidad cognitiva, con complexipy;
  - cobertura de sentencias de las funciones del requisito, a partir de un
    archivo de coverage.py generado con las pruebas de Django.

Uso (desde la raíz del repositorio, con MongoDB levantado):
    COVERAGE_FILE=.coverage coverage run --source=backend/api \
        backend/manage.py test regression_testing.backend
    COBERTURA=.coverage python Metricas/metricas_matriz.py

Igual que metricas_camila.py, trabaja en una base aparte (gym_udem_metricas)
que se borra al terminar.
"""
import os
import sys
from datetime import timedelta
from pathlib import Path

import complexipy
import openpyxl
from radon.complexity import cc_visit

sys.path.insert(0, str(Path(__file__).resolve().parent))
import metricas_camila as base  # noqa: E402  (configura Django y la base de medición)

from api import attendance, features, views  # noqa: E402
from api import db as db_module  # noqa: E402

RAIZ = base.RAIZ
factory = base.factory
ESTUDIANTE, ENTRENADOR, ADMIN = base.ESTUDIANTE, base.ENTRENADOR, base.ADMIN
DOC_ESTUDIANTE = base.DOC_ESTUDIANTE

ORIGEN = RAIZ / "Metricas antiguas.xlsx"
DESTINO = RAIZ / "Metricas nuevas.xlsx"


# ============================================================
# REQUISITOS, FUNCIONES MEDIDAS Y DEFECTOS
# ============================================================

# (requisito, nombre corto de la hoja, funcionalidad, [(archivo, función)])
REQUISITOS = [
    ("RF01", "Registro", "Registrar una cuenta",
     [("views.py", "registrar_cuenta"), ("db.py", "role_for_email")]),
    ("RF02", "Iniciar sesion", "Iniciar sesión",
     [("views.py", "iniciar_sesion")]),
    ("RF03", "Perfil estudiante", "Consultar y actualizar mi perfil",
     [("features.py", "consultar_actualizar_perfil")]),
    ("RF04", "Perfil entrenador", "Consultar el perfil del entrenador",
     [("features.py", "consultar_entrenador"), ("features.py", "_perfil_de_rol")]),
    ("RF05", "Perfil administrador", "Consultar el perfil del administrador",
     [("features.py", "consultar_administrador"), ("features.py", "_perfil_de_rol")]),
    ("RF06", "Bloques y cupos", "Consultar los bloques horarios con sus cupos",
     [("views.py", "consultar_horarios")]),
    ("RF07", "Reservar", "Reservar un bloque para el día siguiente",
     [("views.py", "reservar_mañana")]),
    ("RF08", "Mis reservas", "Consultar mis reservas",
     [("views.py", "consultar_reserva")]),
    ("RF09", "Cancelar reserva", "Cancelar mi reserva",
     [("views.py", "cancelar_reserva")]),
    ("RF10", "Buscar reserva", "Buscar la reserva de un estudiante por su documento",
     [("attendance.py", "buscar_reserva")]),
    ("RF11", "Registrar asistencia", "Registrar la asistencia de un estudiante",
     [("attendance.py", "registrar_asistencia")]),
    ("RF12", "Sin asistencia", "Consultar las reservas sin asistencia registrada",
     [("attendance.py", "consultar_reservas")]),
    ("RF13", "Cerrar jornada", "Procesar las inasistencias al cerrar la jornada",
     [("attendance.py", "procesar_inasistencia"), ("attendance.py", "aplicar_inasistencia")]),
    ("RF14", "Historial", "Ver mi historial",
     [("features.py", "ver_historial")]),
    ("RF15", "Reporte inasistencias", "Ver mi reporte de inasistencias",
     [("attendance.py", "ver_inasistencias")]),
    ("RF16", "Registro diario", "Ver el registro diario del gimnasio (entrenador)",
     [("attendance.py", "ver_registro_entrenador"), ("attendance.py", "_registro_diario_para"),
      ("attendance.py", "build_daily_report")]),
    ("RF20", "Reportar falla", "Reportar una falla o enviar una sugerencia",
     [("features.py", "fallo_sugerencia")]),
    ("RF21", "Buzon", "Consultar el buzón de sugerencias",
     [("features.py", "consultar_buzon"), ("features.py", "eliminar_sugerencia")]),
]

# Defectos encontrados por las pruebas antes de la entrega; ninguno se encontró
# después, así que D = 0 en la EED. Fuentes: INFORME_SEGURIDAD_RENDIMIENTO.md,
# REFACTORIZACION.md, cypress/e2e/defectos/, cypress/component/defectos/ y la
# columna "ID Defecto" de "Casos y Escenarios de Pruebas nuevos.xlsx".
DEFECTOS = {
    "RF01": ["H2 correo con dos arrobas"],
    "RF02": ["VUL-02 sesión sin credenciales", "DEF-03 F5 cierra la sesión"],
    "RF09": ["VUL-04 IDOR al cancelar"],
    "RF10": ["DCTO-01", "DCTO-02"],
    "RF11": ["VUL-01 registrar asistencia sin autenticación"],
    "RF12": ["DEF-05 error tratado como lista vacía"],
    "RF13": ["DEF-01 cierre con fecha futura", "VUL-01 cerrar jornada sin autenticación"],
    "RF16": ["DEF-06 error mostrado como totales en cero"],
    "RF20": ["VUL-03 / DEF-02b mensajes a nombre de otro"],
    "RF21": ["VUL-01 / DEF-02 buzón sin autenticación", "DEF-04 hora en UTC",
             "DEF-07 error tratado como buzón vacío"],
}


# Trazabilidad hacia las pruebas: reglas de negocio (plan, sección 5.1),
# responsable y pruebas automatizadas de regression_testing/ que cubren cada
# requisito. Los escenarios y casos se leen de "Casos y Escenarios de Pruebas nuevos.xlsx".
CASOS_XLSX = RAIZ / "Casos y Escenarios de Pruebas nuevos.xlsx"

DJANGO = {"Violeta Agudelo": "test_violeta.py, test_api_violeta.py, test_seguridad_violeta.py",
          "Sebastián García": "test_sebas.py",
          "Camila Villero": "test_camila.py"}
VITEST = {"Violeta Agudelo": "login.test.jsx, profileView.test.jsx",
          "Sebastián García": "reservas.test.jsx",
          "Camila Villero": "camila.test.jsx, adminPanel.test.jsx, api.test.js"}

TRAZABILIDAD = {
    # rf: (reglas, responsable, cypress, screenplay + gherkin, stagehand)
    "RF01": ("RN01, RN02", "Violeta Agudelo", "violeta/rf01_registro.cy.js",
             "violeta-e2e.feature @rf01 (2); screenpy test_violeta_e2e.py", "Por completar"),
    "RF02": ("RN02, RN09", "Violeta Agudelo", "violeta/rf02_login.cy.js; componente Login.cy.jsx",
             "violeta-e2e.feature @rf02 (2); screenpy test_violeta_e2e.py", "Por completar"),
    "RF03": ("—", "Violeta Agudelo", "violeta/rf03_perfil_estudiante.cy.js",
             "violeta-e2e.feature @rf03 (2); screenpy test_violeta_e2e.py", "Por completar"),
    "RF04": ("—", "Violeta Agudelo", "violeta/rf04_perfil_entrenador.cy.js",
             "violeta-e2e.feature @rf04 (2); screenpy test_violeta_e2e.py", "Por completar"),
    "RF05": ("—", "Violeta Agudelo", "violeta/rf05_perfil_administrador.cy.js",
             "violeta-e2e.feature @rf05 (2); screenpy test_violeta_e2e.py", "Por completar"),
    "RF06": ("RN03, RN04, RN10", "Sebastián García", "sebas/rf06_horarios.cy.js; rf06_horarios_front.cy.js",
             "gym-slots.feature RF06_C1–C2; screenpy test_gym_slots.py", "Por completar"),
    "RF07": ("RN04, RN05, RN06, RN09, RN10", "Sebastián García", "sebas/rf07_reservas.cy.js; rf07_reservas_front.cy.js",
             "gym-slots.feature RF07_C1–C9; screenpy test_gym_slots.py", "Por completar"),
    "RF08": ("—", "Sebastián García", "sebas/rf08_consultar_reservas.cy.js; rf08_consultar_reservas_front.cy.js",
             "gym-slots.feature RF08_C1–C2; screenpy test_gym_slots.py", "Por completar"),
    "RF09": ("RN07, RN11", "Sebastián García", "sebas/rf09_cancelar_reservas.cy.js; rf09_cancelar_reservas_front.cy.js; "
             "seguridad/vulnerabilidades.cy.js (VUL-04)",
             "gym-slots.feature RF09_C1–C4; screenpy test_gym_slots.py", "Por completar"),
    "RF10": ("RN02, RN12", "Sebastián García", "sebas/rf10_buscar_reserva.cy.js; rf10_buscar_reserva_front.cy.js",
             "gym-slots.feature RF10_C1–C4; screenpy test_gym_slots.py", "Por completar"),
    "RF11": ("RN03, RN12", "Camila Villero", "api/rf11_registrar_asistencia.cy.js; ui/panel_entrenador.cy.js; "
             "componente TrainerPanel.cy.jsx; seguridad/vulnerabilidades.cy.js (VUL-01)",
             "camila-e2e.feature @rf11 (3)", "SH-RF12-RF11, SH-RF11-02 (escritas, sin ejecutar)"),
    "RF12": ("RN08", "Camila Villero", "api/rf12_reservas_pendientes.cy.js; ui/panel_entrenador.cy.js; "
             "componente TrainerPanel.cy.jsx; componente defectos/interfaz.cy.jsx (DEF-05)",
             "camila-e2e.feature @rf12 (3)", "SH-RF12-RF11 (escrita, sin ejecutar)"),
    "RF13": ("RN08, RN09", "Camila Villero", "api/rf13_procesar_inasistencias.cy.js; ui/panel_entrenador.cy.js; "
             "componente TrainerPanel.cy.jsx; defectos/defectos.cy.js (DEF-01)",
             "camila-e2e.feature @rf13 (3)", "SH-RF13 (escrita, sin ejecutar)"),
    "RF14": ("—", "Sebastián García", "sebas/rf14_historial_reservas.cy.js; rf14_historial_reservas_front.cy.js",
             "gym-slots.feature RF14_C1–C3; screenpy test_gym_slots.py", "Por completar"),
    "RF15": ("RN08, RN09", "Violeta Agudelo", "violeta/rf15_inasistencias.cy.js",
             "violeta-e2e.feature @rf15 (3); screenpy test_violeta_e2e.py", "Por completar"),
    "RF16": ("RN03, RN08", "Camila Villero", "api/rf16_registro_diario.cy.js; ui/panel_entrenador.cy.js; "
             "componente TrainerPanel.cy.jsx; componente defectos/interfaz.cy.jsx (DEF-06)",
             "camila-e2e.feature @rf16 (3)", "SH-RF12-RF11, SH-RF13 (escritas, sin ejecutar)"),
    "RF20": ("RN11", "Camila Villero", "api/rf20_rf21_buzon.cy.js; ui/buzon.cy.js; componente HistoryView.cy.jsx; "
             "defectos/defectos.cy.js (DEF-02b)",
             "camila-e2e.feature @rf20 (3)", "SH-RF20 (escrita, sin ejecutar)"),
    "RF21": ("—", "Camila Villero", "api/rf20_rf21_buzon.cy.js; ui/buzon.cy.js; componente AdminPanelBuzon.cy.jsx; "
             "defectos (DEF-02, DEF-04, DEF-07)",
             "camila-e2e.feature @rf21 (2)", "SH-RF21 (escrita, sin ejecutar)"),
}


def casos_por_requisito():
    """{rf: (escenarios, casos, probados, pendientes)} desde la hoja "Matriz Casos de Prueba"."""
    if not CASOS_XLSX.exists():
        return {}
    ws = openpyxl.load_workbook(CASOS_XLSX, read_only=True, data_only=True)["Matriz Casos de Prueba"]
    datos = {}
    for fila in ws.iter_rows(min_row=2, values_only=True):
        caso, escenario, rf, estado = fila[0], fila[1], fila[2], fila[10]
        if not caso or not rf:
            continue
        rf = str(rf).strip()
        esc, cps, estados = datos.setdefault(rf, ([], [], []))
        if escenario and str(escenario).replace(" ", "") not in esc:
            esc.append(str(escenario).replace(" ", ""))
        cps.append(str(caso).replace(" ", ""))
        estados.append(str(estado or "").strip())
    return {rf: (", ".join(e), ", ".join(c), est.count("Probado"), len(est) - est.count("Probado"))
            for rf, (e, c, est) in datos.items()}


def hoja_trazabilidad(wb, filas, matriz):
    ws = wb.create_sheet("Trazabilidad requisitos-casos")
    casos = casos_por_requisito()
    encabezado = ["Requisito", "Funcionalidad", "Reglas de negocio", "Responsable", "Escenarios (EP)",
                  "Casos de prueba (CP)", "Estado de los casos", "Cypress", "Screenplay + Gherkin",
                  "Unitarias (Django / Vitest)", "Stagehand", "Defectos encontrados"]
    anchos = [11, 34, 18, 18, 30, 34, 18, 48, 40, 42, 30, 40]
    ws["A1"] = "Matriz de trazabilidad: requisitos, casos y pruebas"
    ws["A1"]._style = matriz["A1"]._style
    for j, (texto, ancho) in enumerate(zip(encabezado, anchos), start=1):
        celda = ws.cell(row=3, column=j, value=texto)
        celda._style = matriz.cell(row=3, column=1)._style
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = ancho
    for i, f in enumerate(filas, start=4):
        reglas, responsable, cypress, screenplay, stagehand = TRAZABILIDAD[f["rf"]]
        esc, cps, probados, pendientes = casos.get(f["rf"], ("Sin escenarios", "Sin casos en el Excel", 0, 0))
        estado = f"{probados} probados, {pendientes} pendientes" if probados + pendientes else "—"
        valores = [f["rf"], f["funcionalidad"], reglas, responsable, esc, cps, estado, cypress, screenplay,
                   f"{DJANGO[responsable]} / {VITEST[responsable]}", stagehand,
                   "; ".join(f["defectos"]) or "Ninguno"]
        for j, valor in enumerate(valores, start=1):
            celda = ws.cell(row=i, column=j, value=valor)
            celda._style = matriz.cell(row=4, column=2)._style
            celda.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B4"


# ============================================================
# MÉTRICAS ESTÁTICAS
# ============================================================

def _ruta(archivo):
    return RAIZ / "backend" / "api" / archivo


def metricas_estaticas(funciones, cobertura):
    """V(G), complejidad cognitiva, líneas y cobertura de las funciones."""
    vg, cognitiva, lineas, ejecutadas, sentencias = [], [], 0, 0, 0
    for archivo, nombre in funciones:
        ruta = _ruta(archivo)
        radon = next(b for b in cc_visit(ruta.read_text()) if b.name == nombre)
        cogn = next(f for f in complexipy.file_complexity(str(ruta)).functions if f.name == nombre)
        vg.append((nombre, radon.complexity))
        cognitiva.append((nombre, cogn.complexity))
        lineas += radon.endline - radon.lineno + 1
        if cobertura:
            todas, faltan = cobertura[str(ruta)]
            rango = range(radon.lineno, radon.endline + 1)
            propias = [n for n in todas if n in rango]
            sentencias += len(propias)
            ejecutadas += len([n for n in propias if n not in faltan])
    return {"vg": vg, "cognitiva": cognitiva, "kloc": lineas / 1000,
            "ejecutadas": ejecutadas, "sentencias": sentencias}


def cargar_cobertura():
    """{ruta absoluta: (sentencias, faltantes)} desde el archivo de coverage.py."""
    archivo = os.getenv("COBERTURA")
    if not archivo:
        print("Sin COBERTURA: la columna de cobertura queda vacía.")
        return None
    import coverage
    cov = coverage.Coverage(data_file=archivo)
    cov.load()
    resultado = {}
    for ruta in cov.get_data().measured_files():
        _, todas, _, faltan, _ = cov.analysis2(ruta)
        resultado[str(Path(ruta).resolve())] = (set(todas), set(faltan))
    return {str(_ruta(a)): resultado[str(_ruta(a).resolve())]
            for a in ("views.py", "features.py", "attendance.py", "db.py")}


# ============================================================
# MEDICIÓN DE LOS REQUISITOS DE VIOLETA Y SEBASTIÁN
# ============================================================

def exigir_2xx(requisito, resultado):
    codigo = resultado[2]
    if not 200 <= codigo < 300:
        sys.exit(f"{requisito}: la vista respondió HTTP {codigo}; la medición no es válida.")
    return resultado


def medir_y_registrar(requisito, descripcion, funcion, preparar, ejecutar, despues=None):
    resultado = exigir_2xx(requisito, base.medir(preparar, ejecutar, despues))
    return base.imprimir(requisito, descripcion, funcion, "-", *resultado)


def manana():
    return db_module.fecha_reserva().isoformat()


def reserva_de_manana(db, slot=1):
    hour = next(inicio for s, inicio, _ in db_module.BLOQUES_HORARIOS if s == slot)
    return db.reservations.insert_one({
        "email": ESTUDIANTE, "slotId": slot, "hour": hour, "reserva_date": manana(),
        "date": db_module.formato_fecha_es(db_module.fecha_reserva()), "estado": "ACTIVA",
        "created_at": db_module.ahora_utc(),
    }).inserted_id


def restablecer_cupos(db):
    db.disponibilidad.update_many({}, [{"$set": {"cupos_disponibles": "$aforo_maximo"}}])


def medir_rf01(db):
    contador = {"n": 0}
    nuevo = {}

    def preparar():
        contador["n"] += 1
        nuevo["email"] = f"metricas.registro{contador['n']}@soyudemedellin.edu.co"
        return factory.post("/api/auth/register/", {
            "name": "Registro Métricas", "email": nuevo["email"],
            "documento": f"98{contador['n']:08d}"}, format="json")

    return medir_y_registrar("RF01", "Registrar una cuenta", "registrar_cuenta (backend/api/views.py)",
                             preparar, views.registrar_cuenta,
                             lambda: db.users.delete_one({"email": nuevo["email"]}))


def medir_rf02(db):
    return medir_y_registrar(
        "RF02", "Iniciar sesión", "iniciar_sesion (backend/api/views.py)",
        lambda: factory.post("/api/auth/login/", {"email": ESTUDIANTE, "documento": DOC_ESTUDIANTE},
                             format="json"),
        views.iniciar_sesion)


def medir_rf03(db):
    return medir_y_registrar(
        "RF03", "Consultar mi perfil", "consultar_actualizar_perfil (backend/api/features.py)",
        lambda: factory.get("/api/users/profile/", {"email": ESTUDIANTE}),
        features.consultar_actualizar_perfil)


def medir_rf04(db):
    return medir_y_registrar(
        "RF04", "Consultar el perfil del entrenador", "consultar_entrenador (backend/api/features.py)",
        lambda: factory.get("/api/users/entrenador/", {"email": ENTRENADOR}),
        features.consultar_entrenador)


def medir_rf05(db):
    return medir_y_registrar(
        "RF05", "Consultar el perfil del administrador", "consultar_administrador (backend/api/features.py)",
        lambda: factory.get("/api/users/administrador/", {"email": ADMIN}),
        features.consultar_administrador)


def medir_rf06(db):
    return medir_y_registrar(
        "RF06", "Consultar los bloques horarios con sus cupos", "consultar_horarios (backend/api/views.py)",
        lambda: factory.get("/api/slots/", {"email": ESTUDIANTE}),
        views.consultar_horarios)


def medir_rf07(db):
    def despues():
        db.reservations.delete_many({"email": ESTUDIANTE})
        restablecer_cupos(db)

    return medir_y_registrar(
        "RF07", "Reservar un bloque para el día siguiente", "reservar_mañana (backend/api/views.py)",
        lambda: factory.post("/api/reservations/", {"email": ESTUDIANTE, "slotId": 1}, format="json"),
        views.reservations, despues)


def medir_rf08(db):
    rid = reserva_de_manana(db)
    resultado = medir_y_registrar(
        "RF08", "Consultar mis reservas", "consultar_reserva (backend/api/views.py)",
        lambda: factory.get("/api/reservations/", {"email": ESTUDIANTE}),
        views.reservations)
    db.reservations.delete_one({"_id": rid})
    return resultado


def medir_rf09(db):
    objetivo = {}

    def preparar():
        objetivo["id"] = reserva_de_manana(db)
        return factory.delete(f"/api/reservations/{objetivo['id']}/")

    def despues():
        db.reservations.delete_many({"email": ESTUDIANTE})
        restablecer_cupos(db)

    db_module.asegurar_disponibilidad(manana())
    return medir_y_registrar(
        "RF09", "Cancelar mi reserva", "cancelar_reserva (backend/api/views.py)",
        preparar, lambda request: views.cancelar_reserva(request, str(objetivo["id"])), despues)


def medir_rf10(db):
    rid = base.nueva_reserva(db, base.hora_de_bloque_iniciado())
    resultado = medir_y_registrar(
        "RF10", "Buscar la reserva de un estudiante por su documento",
        "buscar_reserva (backend/api/attendance.py)",
        lambda: factory.get("/api/students/lookup/", {"documento": DOC_ESTUDIANTE, "actor_email": ENTRENADOR}),
        attendance.buscar_reserva)
    db.reservations.delete_one({"_id": rid})
    return resultado


def medir_rf14(db):
    ayer = (db_module.hoy_local() - timedelta(days=1)).isoformat()
    ids = db.reservations.insert_many([
        {"email": ESTUDIANTE, "slotId": s, "hour": h, "reserva_date": ayer, "date": ayer,
         "estado": e, "created_at": db_module.ahora_utc()}
        for (s, h, _), e in zip(db_module.BLOQUES_HORARIOS, ("COMPLETADA", "NO_SHOW", "CANCELADA"))
    ]).inserted_ids
    resultado = medir_y_registrar(
        "RF14", "Ver mi historial", "ver_historial (backend/api/features.py)",
        lambda: factory.get("/api/reservations/history/", {"email": ESTUDIANTE}),
        features.ver_historial)
    db.reservations.delete_many({"_id": {"$in": ids}})
    return resultado


def medir_rf15(db):
    return medir_y_registrar(
        "RF15", "Ver mi reporte de inasistencias", "ver_inasistencias (backend/api/attendance.py)",
        lambda: factory.get("/api/reports/personal/", {"email": ESTUDIANTE}),
        attendance.ver_inasistencias)


# ============================================================
# EXCEL
# ============================================================

def formula_vg(vg):
    partes = [f"({v - 1}+1)" for _, v in vg]
    return "=" + "+".join(partes) if len(partes) > 1 else f"={vg[0][1] - 1}+1"


def nivel_cognitivo(total):
    return "baja" if total <= 5 else "moderada" if total <= 15 else "alta"


def escribir_excel(filas):
    wb = openpyxl.load_workbook(ORIGEN)
    plantilla = wb["RF01 - Registro"]
    matriz = wb["Matriz de trazabilidad"]
    # Las hojas de la versión anterior usan la numeración vieja: se reemplazan.
    anteriores = [w for w in wb.worksheets if w is not matriz]

    nuevas = []
    for f in filas:
        ws = wb.copy_worksheet(plantilla)
        nuevas.append((ws, f"{f['rf']} - {f['hoja']}"[:31]))
        ws["A1"] = f"{f['rf']} - {f['funcionalidad']}"
        ws["D6"] = formula_vg(f["vg"])
        total_cogn = sum(c for _, c in f["cognitiva"])
        detalle = " + ".join(f"{c} ({n})" for n, c in f["cognitiva"]) if len(f["cognitiva"]) > 1 else ""
        ws["D7"] = f"{detalle + ' = ' if detalle else ''}{total_cogn} ({nivel_cognitivo(total_cogn)})"
        ws["D10"] = f"=({f['ejecutadas']}/{f['sentencias']})*100" if f["sentencias"] else "N/A"
        ws["D11"] = f"={len(f['defectos'])}/({len(f['defectos'])}+0)" if f["defectos"] else "N/A"
        ws["D14"] = f"={len(f['defectos'])}/{f['kloc']:.3f}"
        ws["D15"] = round(f["tiempo_ms"] / 1000, 6)
        ws["D16"] = round(f["cpu"], 2)
        ws["D17"] = round(f["memoria"], 4)
        ws["E14"] = "; ".join(f["defectos"]) if f["defectos"] else None
        ws["E15"] = "segundos (promedio de 20 peticiones)"
        ws["E16"] = "%"
        ws["E17"] = "%"

    for ws in anteriores:
        wb.remove(ws)
    for ws, titulo in nuevas:
        ws.title = titulo
    wb.move_sheet(matriz, offset=len(wb.worksheets) - 1 - wb.index(matriz))

    for fila in matriz.iter_rows(min_row=4, max_row=matriz.max_row):
        for celda in fila:
            celda.value = None
    estilo = [c for c in next(matriz.iter_rows(min_row=4, max_row=4))]
    for i, f in enumerate(filas, start=4):
        valores = [
            f["rf"], f["funcionalidad"], sum(v for _, v in f["vg"]),
            f["ejecutadas"] / f["sentencias"] if f["sentencias"] else "N/A",
            len(f["defectos"]) / f["kloc"],
            1 if f["defectos"] else "N/A",
            f["tiempo_ms"] / 1000, f["cpu"] / 100, f["memoria"] / 100,
        ]
        for j, valor in enumerate(valores):
            celda = matriz.cell(row=i, column=j + 1, value=valor)
            celda._style = estilo[j]._style
        matriz.cell(row=i, column=4).number_format = "0.00%"
        matriz.cell(row=i, column=7).number_format = "0.000000"
    ultima = 3 + len(filas)
    for j, letra in enumerate("CDEFGHI", start=3):
        celda = matriz.cell(row=ultima + 1, column=j, value=f"=AVERAGE({letra}4:{letra}{ultima})")
        celda.number_format = matriz.cell(row=4, column=j).number_format
    matriz["D3"] = "Cobertura estructural (sentencias)"
    matriz["G3"] = "Tiempo de respuesta (s)"
    hoja_trazabilidad(wb, filas, matriz)
    wb.save(DESTINO)


# ============================================================
# EJECUCIÓN
# ============================================================

MEDICIONES = {
    "RF01": medir_rf01, "RF02": medir_rf02, "RF03": medir_rf03, "RF04": medir_rf04,
    "RF05": medir_rf05, "RF06": medir_rf06, "RF07": medir_rf07, "RF08": medir_rf08,
    "RF09": medir_rf09, "RF10": medir_rf10, "RF11": base.medir_rf11, "RF12": base.medir_rf12,
    "RF13": base.medir_rf13, "RF14": medir_rf14, "RF15": medir_rf15, "RF16": base.medir_rf16,
    "RF20": base.medir_rf20,
}


def medir_rf21(db):
    """RF21 tiene dos vistas: se promedian la consulta y el borrado."""
    a, b = base.medir_rf21_consultar(db), base.medir_rf21_eliminar(db)
    return {k: (a[k] + b[k]) / 2 for k in ("promedio_ms", "cpu", "memoria")}


if __name__ == "__main__":
    db = base.conectar()
    cobertura = cargar_cobertura()
    db.client.drop_database(base.settings.MONGO_DB)
    base.preparar_base(db)
    filas = []
    try:
        for rf, hoja, funcionalidad, funciones in REQUISITOS:
            medido = (medir_rf21 if rf == "RF21" else MEDICIONES[rf])(db)
            filas.append({"rf": rf, "hoja": hoja, "funcionalidad": funcionalidad,
                          "defectos": DEFECTOS.get(rf, []),
                          "tiempo_ms": medido["promedio_ms"], "cpu": medido["cpu"],
                          "memoria": medido["memoria"],
                          **metricas_estaticas(funciones, cobertura)})
    finally:
        db.client.drop_database(base.settings.MONGO_DB)

    escribir_excel(filas)
    print(f"\n{'Req.':<6}{'V(G)':>5}{'Cobert.':>9}{'Defect.':>8}{'ms':>9}{'CPU %':>8}{'Mem %':>7}")
    for f in filas:
        cob = f"{f['ejecutadas'] / f['sentencias'] * 100:.1f}" if f["sentencias"] else "-"
        print(f"{f['rf']:<6}{sum(v for _, v in f['vg']):>5}{cob:>9}{len(f['defectos']):>8}"
              f"{f['tiempo_ms']:>9.3f}{f['cpu']:>8.2f}{f['memoria']:>7.2f}")
    print(f"\nExcel escrito en {DESTINO}")
